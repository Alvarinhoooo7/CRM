from docx import Document
from docx.shared import Pt, Inches

doc = Document()
doc.add_heading('Informe Ejecutivo: Logística de Instalación de 33 Equipos', 0)

doc.add_heading('1. Introducción', level=1)
doc.add_paragraph('El presente informe detalla el proceso logístico y la asignación de recursos para la instalación de 33 equipos en 16 localidades a lo largo del país, desde Arica hasta Punta Arenas.')

doc.add_heading('2. Metodología y Herramientas', level=1)
doc.add_paragraph('Para abordar este desafío se desarrolló un CRM personalizado basado en Google Apps Script y Google Sheets. La plataforma integra la API de Google Maps para el cálculo de distancias, optimización de rutas y cálculo de tiempos de traslado reales.')

doc.add_heading('3. Análisis de Rutas y Peajes', level=1)
doc.add_paragraph('El motor logístico prioriza el cumplimiento de la jornada laboral de 9 horas. Para rutas interurbanas (ej. Santiago-Rancagua o Santiago-Valparaíso), se evaluó el costo de peajes frente al costo de horas extra, determinando que el uso de autopistas concesionadas reduce el tiempo de viaje significativamente, evitando sobrecostos laborales.')

doc.add_heading('4. Asignación de Flota y Dotación', level=1)
doc.add_paragraph('Se estandarizó el uso de camionetas para traslados terrestres de los técnicos y herramientas. Para las localidades extremas (Arica, Iquique, Antofagasta, Puerto Montt, Punta Arenas), el sistema permite comparar el costo de vuelos comerciales versus transporte terrestre, recomendando vuelos por el ahorro dramático de días-hombre.')

doc.add_heading('5. Conclusión', level=1)
doc.add_paragraph('La automatización del cálculo logístico redujo el tiempo de planificación a segundos, asegurando trazabilidad financiera y eficiencia operativa. Todos los costos de viáticos, alojamiento y pasajes quedaron justificados algorítmicamente.')

doc.save('Informe_Logistica.docx')
print("Docx generado correctamente.")

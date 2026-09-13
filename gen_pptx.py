from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

# Slide 1: Title
slide_layout = prs.slide_layouts[0] 
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Despliegue Logístico a Nivel Nacional"
subtitle.text = "Instalación de 33 Equipos en 16 Localidades\nTecnologías Aplicadas a los Sistemas Inteligentes"

# Slide 2: El Desafío
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "El Desafío Operativo"
content = slide.placeholders[1].text_frame
content.text = "Desplegar cuadrillas técnicas a lo largo de Chile optimizando recursos."
p = content.add_paragraph()
p.text = "33 equipos a instalar en 16 localidades."
p = content.add_paragraph()
p.text = "Variabilidad geográfica: Desde Arica hasta Punta Arenas."
p = content.add_paragraph()
p.text = "Control estricto de presupuesto, horas extra y transporte."

# Slide 3: La Solución (CRM)
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Solución Inteligente: CRM Logístico"
content = slide.placeholders[1].text_frame
content.text = "Desarrollo de un sistema en Google Apps Script + Sheets."
p = content.add_paragraph()
p.text = "Integración nativa con Google Maps API para cálculo de rutas y tiempos."
p = content.add_paragraph()
p.text = "Interfaz de usuario tipo 'Bitácora' para registro rápido."
p = content.add_paragraph()
p.text = "Motor de decisiones que sugiere transporte (Terrestre vs Vuelo) basado en Costo de Oportunidad."

# Slide 4: Resultados
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Impacto y Resultados"
content = slide.placeholders[1].text_frame
content.text = "Reducción dramática del tiempo de planificación a segundos."
p = content.add_paragraph()
p.text = "Disminución de costos al seleccionar autopistas de forma inteligente (evitando horas extra)."
p = content.add_paragraph()
p.text = "Visualización consolidada de métricas operacionales."

prs.save('Presentacion_Logistica.pptx')
print("PPTX generado correctamente.")
